#!/usr/bin/env python3
"""Build an editable Word version of the stage-1 report from docs/REPORT_STAGE1.md.

The report is authored in Markdown, so generating DOCX from the source keeps headings,
tables and lists intact and editable — unlike scraping text back out of the PDF.

Requires python-docx. The project's conda environment deliberately does not ship it (the app
has no Word dependency); the DSH bundled Python does, e.g.:

    & "$env:DSH_HOME\\..\\dsh-runtimes\\dsh-primary-runtime\\dependencies\\python\\python.exe" `
        scripts/build_report_docx.py --name 姓名 --student-id 学号

    python scripts/build_report_docx.py                     # -> dist/Lab4-stage1-report.docx

Cover name/student-id default to blank lines; values are never stored in this file.
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "docs" / "REPORT_STAGE1.md"
DEFAULT_OUTPUT = ROOT / "dist" / "Lab4-stage1-report.docx"

FONT = "Microsoft YaHei"
MONO = "Consolas"
ACCENT = RGBColor(0x14, 0x51, 0x4A)
MUTED = RGBColor(0x56, 0x61, 0x5E)


# ---------------------------------------------------------------- markdown parsing

def inline_runs(text: str) -> list[tuple[str, str]]:
    """Split inline Markdown into (kind, text) runs: plain / bold / code."""
    text = re.sub(r"\[([^\]]+)\]\(([^)\s]+)\)", r"\1", text)  # internal references print as their file name
    runs: list[tuple[str, str]] = []
    for token in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            runs.append(("bold", token[2:-2]))
        elif token.startswith("`") and token.endswith("`"):
            runs.append(("code", token[1:-1]))
        else:
            runs.append(("plain", token))
    return runs


def parse(markdown: str) -> list[dict]:
    lines = markdown.splitlines()
    blocks: list[dict] = []
    index = 0
    while index < len(lines):
        stripped = lines[index].strip()

        if stripped.startswith("```"):
            index += 1
            body: list[str] = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                body.append(lines[index])
                index += 1
            blocks.append({"type": "code", "text": "\n".join(body)})
            index += 1
            continue

        if re.match(r"^#{1,4} ", stripped):
            level = len(stripped) - len(stripped.lstrip("#"))
            blocks.append({"type": f"h{level}", "text": stripped[level + 1:].strip()})
            index += 1
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|$", lines[index + 1].strip()):
            header = [cell.strip() for cell in stripped.strip("|").split("|")]
            index += 2
            rows: list[list[str]] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append([cell.strip() for cell in lines[index].strip().strip("|").split("|")])
                index += 1
            blocks.append({"type": "table", "header": header, "rows": rows})
            continue

        if re.match(r"^\s*[-*] ", lines[index]):
            items = []
            while index < len(lines) and re.match(r"^\s*[-*] ", lines[index]):
                items.append(re.sub(r"^\s*[-*] ", "", lines[index]).strip())
                index += 1
            blocks.append({"type": "ul", "items": items})
            continue

        if re.match(r"^\s*\d+\. ", lines[index]):
            items = []
            while index < len(lines) and re.match(r"^\s*\d+\. ", lines[index]):
                items.append(re.sub(r"^\s*\d+\. ", "", lines[index]).strip())
                index += 1
            blocks.append({"type": "ol", "items": items})
            continue

        if stripped.startswith("> "):
            blocks.append({"type": "quote", "text": stripped[2:].strip()})
            index += 1
            continue

        if stripped in {"---", "***", "___"}:
            index += 1
            continue

        if not stripped:
            index += 1
            continue

        paragraph: list[str] = []
        while index < len(lines) and lines[index].strip() and not re.match(
                r"^(#{1,4} |\s*[-*] |\s*\d+\. |\||> |```|---$)", lines[index].strip()):
            paragraph.append(lines[index].strip())
            index += 1
        blocks.append({"type": "p", "text": " ".join(paragraph)})
    return blocks


# ---------------------------------------------------------------- docx helpers

def set_style_font(style, size: float, bold: bool = False, mono: bool = False) -> None:
    style.font.name = MONO if mono else FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.element.rPr.rFonts.set(qn("w:eastAsia"), MONO if mono else FONT)


def shade(cell, colour: str) -> None:
    element = OxmlElement("w:shd")
    element.set(qn("w:val"), "clear")
    element.set(qn("w:fill"), colour)
    cell._tc.get_or_add_tcPr().append(element)


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin, separate, end = (OxmlElement("w:fldChar") for _ in range(3))
    begin.set(qn("w:fldCharType"), "begin")
    separate.set(qn("w:fldCharType"), "separate")
    end.set(qn("w:fldCharType"), "end")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instruction
    run._r.append(begin)
    run._r.append(text)
    run._r.append(separate)
    run._r.append(end)


def write_runs(paragraph, text: str) -> None:
    for kind, value in inline_runs(text):
        run = paragraph.add_run(value)
        if kind == "bold":
            run.bold = True
        elif kind == "code":
            run.font.name = MONO
            run.font.size = Pt(9.5)
            run._element.rPr.rFonts.set(qn("w:eastAsia"), MONO)


def build(markdown: str, name: str, student_id: str) -> Document:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    for attr, value in (("top_margin", 2.2), ("bottom_margin", 2.2), ("left_margin", 2.0), ("right_margin", 2.0)):
        setattr(section, attr, Cm(value))

    normal = document.styles["Normal"]
    set_style_font(normal, 10.5)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.3
    grid = section._sectPr.find(qn("w:docGrid"))
    if grid is not None:  # the default template's line grid forces 18pt lines on CJK text
        grid.set(qn("w:type"), "none")
        grid.attrib.pop(qn("w:linePitch"), None)
    for level, size in ((1, 15), (2, 12.5), (3, 11)):
        style = document.styles[f"Heading {level}"]
        set_style_font(style, size, bold=True)
        style.font.color.rgb = RGBColor(0x1B, 0x24, 0x22)
        style.paragraph_format.space_before = Pt(10 if level == 1 else 8)
        style.paragraph_format.space_after = Pt(4)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("MongoDB 在线问卷系统 · 第一阶段项目完成报告  ·  第 ")
    add_field(footer, "PAGE")
    footer.add_run(" 页")
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = MUTED

    # cover
    title = document.add_paragraph()
    title.paragraph_format.space_before = Pt(120)
    run = title.add_run("MongoDB 在线问卷系统")
    run.font.size = Pt(26)
    run.bold = True
    subtitle = document.add_paragraph()
    run = subtitle.add_run("第一阶段项目完成报告")
    run.font.size = Pt(26)
    run.bold = True
    meta = document.add_paragraph()
    meta.paragraph_format.space_before = Pt(30)
    for label, value in (("课程：", "数据库管理系统"), ("项目：", "基于 MongoDB 的问卷系统开发（第一阶段）"),
                         ("仓库：", "https://github.com/oceancat91/DB-Lab.git"),
                         ("姓名：", name or "＿＿＿＿＿＿＿＿"), ("学号：", student_id or "＿＿＿＿＿＿＿＿"),
                         ("日期：", "2026 年 10 月 16 日")):
        line = document.add_paragraph()
        line.paragraph_format.space_after = Pt(2)
        head = line.add_run(label)
        head.bold = True
        head.font.color.rgb = MUTED
        line.add_run(value)
    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(24)
    run = note.add_run("本报告对应《大作业一第一阶段》“七、提交内容”第 4 项，覆盖项目目标、系统设计、"
                       "MongoDB 设计、跳转设计、校验设计、AI 使用过程、测试结果、遇到的问题及解决办法。")
    run.font.size = Pt(9.5)
    run.font.color.rgb = MUTED
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    blocks = parse(markdown)
    headings = [block["text"] for block in blocks if block["type"] == "h2" and block["text"] != "目录"]

    toc_heading = document.add_paragraph()
    run = toc_heading.add_run("目录")
    run.bold = True
    run.font.size = Pt(15)
    for text in headings:
        line = document.add_paragraph()
        line.paragraph_format.space_after = Pt(2)
        line.paragraph_format.left_indent = Cm(0.6)
        line.add_run(text)
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    first_h1 = True
    for block in blocks:
        kind = block["type"]
        if kind == "h1":
            if first_h1:  # the cover already carries the title
                first_h1 = False
                continue
            document.add_paragraph(block["text"], style="Heading 1")
        elif kind == "h2":
            document.add_paragraph(block["text"], style="Heading 1")
        elif kind == "h3":
            document.add_paragraph(block["text"], style="Heading 2")
        elif kind == "h4":
            document.add_paragraph(block["text"], style="Heading 3")
        elif kind == "p":
            paragraph = document.add_paragraph()
            # Justified CJK stretches badly when a long code token has to move to the next line.
            long_code = any(kind_ == "code" and len(value) >= 12 for kind_, value in inline_runs(block["text"]))
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if long_code else WD_ALIGN_PARAGRAPH.JUSTIFY
            write_runs(paragraph, block["text"])
        elif kind in {"ul", "ol"}:
            style = "List Bullet" if kind == "ul" else "List Number"
            for item in block["items"]:
                paragraph = document.add_paragraph(style=style)
                paragraph.paragraph_format.space_after = Pt(2)
                write_runs(paragraph, item)
        elif kind == "quote":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.6)
            write_runs(paragraph, block["text"])
            for run in paragraph.runs:
                run.font.color.rgb = MUTED
        elif kind == "code":
            paragraph = document.add_paragraph()
            run = paragraph.add_run(block["text"])
            run.font.name = MONO
            run.font.size = Pt(9)
            run._element.rPr.rFonts.set(qn("w:eastAsia"), MONO)
        elif kind == "table":
            table = document.add_table(rows=1, cols=len(block["header"]))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for cell, text in zip(table.rows[0].cells, block["header"]):
                cell.text = ""
                write_runs(cell.paragraphs[0], text)
                for run in cell.paragraphs[0].runs:
                    run.bold = True
                shade(cell, "EEF2EF")
            for row in block["rows"]:
                cells = table.add_row().cells
                for cell, text in zip(cells, row):
                    cell.text = ""
                    write_runs(cell.paragraphs[0], text)
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after = Pt(1)
                        for run in paragraph.runs:
                            run.font.size = Pt(9)
            document.add_paragraph().paragraph_format.space_after = Pt(2)
    return document


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", default=str(SOURCE))
    parser.add_argument("--output", default=str(DEFAULT_OUTPUT))
    parser.add_argument("--name", default="", help="封面姓名；留空则打印空白填写行（不写入仓库）")
    parser.add_argument("--student-id", default="", help="封面学号；留空则打印空白填写行（不写入仓库）")
    arguments = parser.parse_args()

    source, output = Path(arguments.source), Path(arguments.output)
    if not source.is_file():
        print(f"源报告不存在：{source}")
        return 1
    document = build(source.read_text(encoding="utf-8"), arguments.name, arguments.student_id)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(f"源：{source}")
    print(f"DOCX：{output}（{output.stat().st_size} 字节）")
    return 0


if __name__ == "__main__":
    sys.exit(main())
